"""Preenche a planilha Open Cost do BID Alpargatas com o time sugerido (tela 13 do deck).

Custo (colunas F, G e H) fica em branco para o Daniel; I, J e K seguem as fórmulas do modelo.
Horas Mês = FTE x horas/mês de 1 FTE (célula B3 do DTC; o B2B lê a mesma célula).
Saída em proposta/ (fora do repositório público).

Rodar: python3 fonte/preencher_open_cost.py <planilha_modelo.xlsx> <saida.xlsx>
"""
import sys
from copy import copy

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

H_FTE_MES = 176   # 22 dias úteis x 8 h; o deck usa 140 h produtivas = 80% disso

DTC = [
    ('Gestor(a) de conta DTC', 'Sênior', 1, 1.00,
     'Atendimento exclusivo do DTC: pauta, cronograma integrado de 100 dias úteis, SharePoint e status.'),
    ('Designer 3D e engenharia de PDV', 'Pleno', 1, 1.00,
     'Engenharia por formato de loja, projeto executivo e aprovação da peça em AR.'),
    ('Designer gráfico e arte-final', 'Pleno', 1, 1.00,
     'Desdobramento do KV nos formatos de loja e arte-final para a gráfica.'),
    ('Produtor(a) de fornecedores', 'Sênior', 1, 1.00,
     'Concorrência de três orçamentos, gestão do fornecedor e QA de lote.'),
    ('Produtor(a) de campo e SAC', 'Pleno', 1, 1.00,
     'Protótipo testado em loja e atendimento à rede por e-mail e WhatsApp.'),
    ('Analista de logística', 'Pleno', 1, 1.00,
     'Plano de envio loja a loja para as 728 lojas, entrega e ocorrências.'),
    ('Diretor(a) de criação (compartilhado)', 'Sênior', 1, 0.30,
     'Mesma pessoa no B2B (30%). Uma só direção de criação para a unicidade da marca.'),
    ('Arquiteto(a) e 3D (compartilhado)', 'Sênior', 1, 0.25,
     'Mesma pessoa no B2B (50%). Projetos especiais: corner, ilha e cenografia.'),
    ('Analista fiscal (compartilhado)', 'Pleno', 1, 0.25,
     'Mesma pessoa no B2B (25%). Faturamento, SEFAZ e pagamentos.'),
]

B2B = [
    ('Gestor(a) de conta B2B', 'Sênior', 1, 1.00,
     'Canais alimentar e especializado: pauta, cronograma e status.'),
    ('Designer gráfico sênior', 'Sênior', 1, 1.00,
     'KV, desdobramento e as 16 peças de cada campanha.'),
    ('Designer de artes rápidas', 'Pleno', 1, 1.00,
     'As 8 artes rápidas da semana e vídeos de até 30 segundos.'),
    ('Produtor(a) de fornecedores', 'Pleno', 1, 1.00,
     'Orçamento, negociação dos itens estratégicos e amostras.'),
    ('Analista de logística e implementação', 'Pleno', 1, 0.75,
     'Operador logístico ou cliente, ambientações de grandes contas e ativações.'),
    ('Diretor(a) de criação (compartilhado)', 'Sênior', 1, 0.30,
     'Mesma pessoa do DTC (30%). Uma só direção de criação para a unicidade da marca.'),
    ('Arquiteto(a) e 3D (compartilhado)', 'Sênior', 1, 0.50,
     'Mesma pessoa do DTC (25%). Ambientações de grandes contas e projetos especiais.'),
    ('Analista fiscal (compartilhado)', 'Pleno', 1, 0.25,
     'Mesma pessoa do DTC (25%). Faturamento, SEFAZ e pagamentos.'),
]

# hora média por entrega medida no Ariuno (mesmos números da tela 14 do deck)
PROJETOS = [
    ('Arte rápida', 'Post, banner ou peça digital a partir do KV vigente', 'B2B', 2.5,
     'Prazo de 1 a 2 dias úteis. Horas: média medida no sistema da agência.'),
    ('Peça tática', 'Adaptação e arte-final de wobbler, stopper, faixa de gôndola e similares', 'DTC/B2B', 3.2,
     'A partir do KV aprovado. Horas: média medida no sistema da agência.'),
    ('KV de campanha', 'Conceito visual da campanha com desdobramento base', 'DTC/B2B', 10,
     'Horas: média dos KVs com desdobramento medidos no sistema da agência.'),
    ('Peça estrutural', 'Engenharia e projeto executivo de display, quiosque, prateleira ou vitrine', 'DTC/B2B', 8.6,
     'Executivo pronto para produção. Horas: média medida no sistema da agência.'),
    ('Aprovação em 3D e AR', 'Render 3D e visualização em AR por formato de loja', 'DTC', 2.5,
     'Por formato, antes da produção. Horas: mediana medida no sistema da agência.'),
    ('Ambientação de grande conta', 'Projeto de ambientação de loja de grande conta', 'B2B', 19.7,
     'Horas: média dos projetos de display medidos no sistema da agência.'),
    ('Projeto especial', 'Stand, corner, ilha ou cenografia', 'DTC/B2B', 72,
     'Horas: mediana dos stands e cenografias medidos (45 a 137 h por projeto).'),
    ('Produção de item estrutural', 'Orçamento com três fornecedores, protótipo, acompanhamento e QA de lote', 'DTC/B2B', 10,
     'Por item. Orçamento 3,5 h + acompanhamento 4,1 h + 30% de protótipo e QA.'),
    ('Produção de item tático', 'Orçamento e acompanhamento de item repetido do enxoval', 'DTC/B2B', 4,
     'Por item. Horas: orçamento e acompanhamento medidos no sistema da agência.'),
]

AMARELO = PatternFill('solid', fgColor='FFF2CC')
CINZA = Font(name='Segoe UI', size=10, color='FF555555')
FINA = Side(style='thin')


def preenche_time(ws, linhas, base_horas):
    # estilos das linhas modelo (6 branca, 7 listrada, 11 total) antes de sobrescrever
    snap = {r: {c: (copy(x.font), copy(x.fill), copy(x.border), x.number_format, copy(x.alignment))
                for c in 'ABCDEFGHIJKL' for x in [ws[f'{c}{r}']]} for r in (6, 7, 11)}

    def aplica(r, origem):
        for c in 'ABCDEFGHIJKL':
            x = ws[f'{c}{r}']
            x.font, x.fill, x.border, x.number_format, x.alignment = (copy(v) for v in snap[origem][c])

    ws['A3'] = 'Horas/mês de 1 FTE'
    ws['A3'].font = CINZA
    ws['B3'] = base_horas
    ws['B3'].font = Font(name='Segoe UI', size=10, bold=True)
    ws['B3'].number_format = '#,##0" h"'
    ws['B3'].alignment = Alignment(horizontal='center')
    ws['B3'].border = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)
    ws['C3'] = ('22 dias úteis × 8 h. Horas Mês = FTE × esta base. '
                'Custo Salarial/Hora = (salário + encargos) ÷ esta base.')
    ws['C3'].font = CINZA

    ini = 6
    fim = ini + len(linhas) - 1
    tot = fim + 1
    # o bloco do modelo (linhas 6 a 11) traz fórmulas do total antigo: limpa antes de reescrever
    for r in range(ini, max(tot, 11) + 1):
        for c in 'ABCDEFGHIJKL':
            ws[f'{c}{r}'].value = None
    for i, (cargo, sen, qtd, fte, obs) in enumerate(linhas):
        r = ini + i
        aplica(r, 6 if i % 2 == 0 else 7)
        ws[f'A{r}'] = cargo
        ws[f'B{r}'] = sen
        ws[f'C{r}'] = qtd
        ws[f'D{r}'] = fte
        ws[f'E{r}'] = f'=D{r}*$B$3'
        ws[f'I{r}'] = f'=F{r}*(1+G{r})*(1+H{r})'
        ws[f'J{r}'] = f'=I{r}*E{r}'
        ws[f'K{r}'] = f'=C{r}*J{r}'
        ws[f'L{r}'] = obs
        for c in 'AL':
            al = copy(ws[f'{c}{r}'].alignment)
            ws[f'{c}{r}'].alignment = Alignment(horizontal=al.horizontal, vertical='center', wrap_text=True)
        for c in 'BCDEFGHIJK':
            al = copy(ws[f'{c}{r}'].alignment)
            ws[f'{c}{r}'].alignment = Alignment(horizontal=al.horizontal, vertical='center')

    aplica(tot, 11)
    canal = 'DTC' if 'DTC' in ws['A1'].value else 'B2B'
    ws[f'A{tot}'] = f'TOTAL FEE MENSAL {canal}'
    ws[f'C{tot}'] = f'=SUM(C{ini}:C{fim})'
    ws[f'D{tot}'] = f'=SUM(D{ini}:D{fim})'
    ws[f'D{tot}'].number_format = '0.0" FTE"'
    ws[f'E{tot}'] = f'=SUM(E{ini}:E{fim})'
    ws[f'E{tot}'].number_format = '#,##0'
    ws[f'G{tot}'] = f'=IFERROR(AVERAGE(G{ini}:G{fim}),"")'
    ws[f'H{tot}'] = f'=IFERROR(AVERAGE(H{ini}:H{fim}),"")'
    ws[f'K{tot}'] = f'=SUM(K{ini}:K{fim})'
    ws[f'L{tot}'] = ('Headcount soma 1 por linha; os 3 compartilhados aparecem no DTC e no B2B '
                     '(14 pessoas, 12,6 FTE nos dois canais).')
    ws[f'L{tot}'].font = Font(name='Segoe UI', size=9, bold=False, color='FF555555')
    ws[f'L{tot}'].alignment = Alignment(vertical='center', wrap_text=True)
    for c in 'ABCDEFGHIJK':
        ws[f'{c}{tot}'].alignment = Alignment(horizontal=ws[f'{c}{tot}'].alignment.horizontal, vertical='center')

    # amarelo só enquanto o custo estiver vazio: some quando o Daniel preenche
    ws.conditional_formatting.add(f'F{ini}:H{fim}', FormulaRule(formula=[f'ISBLANK(F{ini})'], fill=AMARELO))
    return ini, fim, tot


def preenche_projetos(ws):
    snap = {r: {c: (copy(ws[f'{c}{r}'].font), copy(ws[f'{c}{r}'].fill), copy(ws[f'{c}{r}'].border),
                    ws[f'{c}{r}'].number_format, copy(ws[f'{c}{r}'].alignment)) for c in 'ABCDEFG'}
            for r in (6, 7)}
    for i, (cat, desc, canal, horas, prem) in enumerate(PROJETOS):
        r = 6 + i
        for c in 'ABCDEFG':
            x = ws[f'{c}{r}']
            x.font, x.fill, x.border, x.number_format, x.alignment = (copy(v) for v in snap[6 if i % 2 == 0 else 7][c])
        ws[f'A{r}'] = cat
        ws[f'B{r}'] = desc
        ws[f'C{r}'] = canal
        ws[f'D{r}'] = horas
        ws[f'D{r}'].number_format = '#,##0.0'
        ws[f'G{r}'] = prem
        for c in 'ABG':
            ws[f'{c}{r}'].alignment = Alignment(vertical='center', wrap_text=True)
        for c in 'CDEF':
            ws[f'{c}{r}'].alignment = Alignment(horizontal=ws[f'{c}{r}'].alignment.horizontal, vertical='center')
    fim = 6 + len(PROJETOS) - 1
    ws.conditional_formatting.add(f'E6:F{fim}', FormulaRule(formula=['ISBLANK(E6)'], fill=AMARELO))
    ws['A3'] = ('Horas Estimadas: hora média por entrega medida no sistema de gestão da 75 LAB '
                '(11.060 apontamentos, jul a out/2026).')
    ws['A3'].font = CINZA


def main(modelo, saida):
    wb = openpyxl.load_workbook(modelo)
    dtc = preenche_time(wb['Estrutura DTC'], DTC, H_FTE_MES)
    ws_b2b = wb['Estrutura B2B']
    b2b = preenche_time(ws_b2b, B2B, H_FTE_MES)
    ws_b2b['B3'] = "='Estrutura DTC'!B3"
    preenche_projetos(wb['Proposta Híbrida - Projetos'])
    wb.save(saida)
    print('DTC linhas', dtc, '| B2B linhas', b2b)


if __name__ == '__main__':
    main(*sys.argv[1:3])
